"""Local file input and strict analysis-configuration checks."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pandas as pd


def _check_headers(headers: list) -> None:
    if not headers or any(not isinstance(x, str) or not x.strip() for x in headers):
        raise ValueError("Every input column needs a nonempty text header.")
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column headers are not allowed.")
    if any(x != x.strip() for x in headers):
        raise ValueError("Remove leading or trailing whitespace from column headers.")


def read_data(path: str | Path, subject_col: str | None = None) -> pd.DataFrame:
    """Read CSV or XLSX; do not silently repair duplicated column names.

    Empty cells and the exact strings NA/NaN are missing. Values such as '?'
    or 'not available' are retained so numerical validation can reject them.
    CSV subject identifiers are preserved as strings, including leading zeros.
    """
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"Input file does not exist: {path}")
    suffix = path.suffix.lower()
    dtypes = {subject_col: "string"} if subject_col else None
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, strict=True)
            try:
                headers = next(reader, [])
                _check_headers(headers)
                for row in reader:
                    if row and len(row) != len(headers):
                        raise ValueError(f"CSV record ending on line {reader.line_num} has {len(row)} fields; expected {len(headers)}. Repair the row before analysis.")
            except csv.Error as exc:
                raise ValueError(f"Invalid CSV syntax: {exc}.") from exc
        data = pd.read_csv(path, encoding="utf-8-sig", dtype=dtypes,
                           keep_default_na=False, na_values=["", "NA", "NaN"])
    elif suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("XLSX input requires the excel extra: pip install '.[excel]'.") from exc
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            headers = list(next(book.worksheets[0].iter_rows(values_only=True), []))
            _check_headers(headers)
            for values, original in zip(book.worksheets[0].iter_rows(), formulas.worksheets[0].iter_rows()):
                for cached, source in zip(values, original):
                    if cached.data_type == "e" or source.data_type == "e":
                        raise ValueError(f"Excel error cell {source.coordinate}: {cached.value}. Correct the formula/value before analysis.")
                    if source.data_type == "f" and cached.value is None:
                        raise ValueError(f"Excel formula {source.coordinate} has no cached result. Recalculate and save the workbook in Excel, or export verified values to CSV.")
        finally:
            book.close()
            formulas.close()
        data = pd.read_excel(path, sheet_name=0, engine="openpyxl", dtype=dtypes,
                             keep_default_na=False, na_values=["", "NA", "NaN"])
    else:
        raise ValueError("Input must be a UTF-8 CSV or an XLSX workbook (first worksheet).")
    if data.empty:
        raise ValueError("The input contains no data rows.")
    return data


def _keys(value: dict, allowed: set[str], where: str) -> None:
    if not isinstance(value, dict):
        raise ValueError(f"{where} must be a JSON object.")
    unknown = set(value) - allowed
    if unknown:
        raise ValueError(f"Unknown {where} keys: {sorted(unknown)}. Check for spelling mistakes.")


def _columns(value, label: str, allow_empty: bool = True) -> list[str]:
    if not isinstance(value, list) or any(not isinstance(x, str) or not x.strip() for x in value):
        raise ValueError(f"{label} must be a list of nonempty column names.")
    if len(value) != len(set(value)) or (not allow_empty and not value):
        raise ValueError(f"{label} must contain distinct column names and cannot be empty here.")
    return value


def validate_config(config: dict) -> dict:
    """Validate without guessing misspelled or omitted analysis choices."""
    _keys(config, {"title", "subject_id", "group", "group_order", "continuous", "categorical",
                   "comparisons", "linear", "icc"}, "configuration")
    for name in ["subject_id", "group"]:
        if not isinstance(config.get(name), str) or not config[name].strip():
            raise ValueError(f"A nonempty {name!r} column name is required.")
    if config["subject_id"] == config["group"]:
        raise ValueError("subject_id and group must be different columns.")
    if "title" in config and not isinstance(config["title"], str):
        raise ValueError("title must be text.")
    continuous = _columns(config.get("continuous", []), "continuous")
    categorical = _columns(config.get("categorical", []), "categorical")
    if not continuous and not categorical:
        raise ValueError("Select at least one continuous or categorical variable.")
    if set(continuous) & set(categorical):
        raise ValueError("continuous and categorical variable lists must not overlap.")
    if config["subject_id"] in continuous + categorical:
        raise ValueError("The patient identifier must not be analyzed as a measurement.")
    if config.get("group_order") is not None:
        order = config["group_order"]
        if not isinstance(order, list) or len(order) != 2 or order[0] == order[1]:
            raise ValueError("group_order must contain two distinct observed group labels.")
    comp = config.get("comparisons", {})
    _keys(comp, {"enabled", "continuous_test", "adjust_method"}, "comparisons")
    if "enabled" in comp and not isinstance(comp["enabled"], bool):
        raise ValueError("comparisons.enabled must be true or false, not text.")
    if comp.get("continuous_test", "welch") not in {"welch", "mannwhitney"}:
        raise ValueError("continuous_test must be 'welch' or 'mannwhitney'.")
    if comp.get("adjust_method", "fdr_bh") not in {"fdr_bh", "fdr_by", "none"}:
        raise ValueError("adjust_method must be fdr_bh, fdr_by, or none.")
    if "linear" in config:
        model = config["linear"]
        _keys(model, {"outcome", "predictors", "categorical"}, "linear")
        if not isinstance(model.get("outcome"), str) or not model["outcome"].strip():
            raise ValueError("linear.outcome must name the continuous outcome.")
        _columns(model.get("predictors"), "linear.predictors", allow_empty=False)
        _columns(model.get("categorical", []), "linear.categorical")
    if "icc" in config:
        reliability = config["icc"]
        _keys(reliability, {"raters", "subject_id", "n_boot", "seed"}, "icc")
        _columns(reliability.get("raters"), "icc.raters", allow_empty=False)
        if "subject_id" in reliability and (
            not isinstance(reliability["subject_id"], str) or not reliability["subject_id"].strip()
        ):
            raise ValueError("icc.subject_id must be a nonempty column name.")
    return config


def load_config(path: str | Path) -> dict:
    def no_duplicate_keys(pairs):
        out = {}
        for key, value in pairs:
            if key in out:
                raise ValueError(f"Duplicate JSON key: {key!r}.")
            out[key] = value
        return out

    with Path(path).open(encoding="utf-8-sig") as handle:
        config = json.load(handle, object_pairs_hook=no_duplicate_keys)
    return validate_config(config)


def template_config() -> dict:
    return {
        "title": "Clinical cohort analysis",
        "subject_id": "patient_id",
        "group": "group",
        "group_order": ["Control", "Intervention"],
        "continuous": ["age", "bmi", "oks_pre", "oks_post"],
        "categorical": ["sex"],
        "comparisons": {"enabled": False, "continuous_test": "welch", "adjust_method": "fdr_bh"},
        "linear": {"outcome": "oks_post", "predictors": ["group", "oks_pre", "age", "bmi", "sex"],
                   "categorical": ["group", "sex"]},
    }
